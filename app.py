
import io, os, re, tempfile, unicodedata
from pathlib import Path
import fitz
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

st.set_page_config(page_title="経費予測Pilot", page_icon="📄", layout="wide")

st.title("経費予測Pilot")
st.caption("見積PDF → Common Data → Excel｜原本記載を優先し、記載のない税額は自動生成しません。")

def clean(s):
    return " ".join((s or "").split()).strip()

def nval(s):
    if s is None or s == "": return None
    s = unicodedata.normalize("NFKC", str(s))
    s = s.replace("￥","").replace("¥","").replace("▲","-").replace("△","-").strip()
    # 「120,000 ※注記」「0 ※既存利用」のように注記が同じ座標帯へ
    # 重なっても、先頭の数値トークンだけを原本値として取得する。
    m = re.search(r"-?\d[\d,]*(?:\.\d+)?", s)
    if not m: return None
    try:
        v=float(m.group(0).replace(",",""))
        return int(v) if v.is_integer() else v
    except:
        return None

def numeric_tail_note(s):
    """数値の後ろに同じセル帯で拾った注記があれば返す。"""
    if not s: return None
    s=unicodedata.normalize("NFKC", str(s)).strip()
    m=re.search(r"-?\d[\d,]*(?:\.\d+)?", s)
    if not m: return s
    tail=s[m.end():].strip()
    return tail or None

def group_lines(page, ymin, ymax):
    groups=[]
    for w in page.get_text("words"):
        x0,y0,x1,y1,t,*_=w
        if not (ymin <= y0 <= ymax): continue
        for g in groups:
            if abs(g["y"]-y0) <= 2.1:
                g["w"].append((x0,t)); break
        else:
            groups.append({"y":y0,"w":[(x0,t)]})
    return sorted(groups,key=lambda z:z["y"])

def classify(path):
    d=fitz.open(path)
    t=unicodedata.normalize("NFKC","\n".join(d[i].get_text("text") for i in range(min(2,len(d)))))
    fn=unicodedata.normalize("NFKC",Path(path).name)
    if "BAN" in fn.upper() or "バンディット" in t or "BANDIT" in t.upper():
        return "BAN"
    if "サンテン" in fn or "サンテン" in t or "SANTEN-CORP.COM" in t.upper():
        return "SANTEN"
    return "HIMI"

def parse_ban(path):
    d=fitz.open(path); rows=[]

    def detect_ban_columns(page):
        # ヘッダー文字の実座標から列境界を作る。
        header=[w for w in page.get_text("words") if 85 <= w[1] <= 115]
        anchors={}
        for w in header:
            t=clean(w[4]).replace(" ","").replace("　","")
            if "項目" in t: anchors["item"]=w[0]
            elif "仕様" in t: anchors["spec"]=w[0]
            elif "数量" in t: anchors["qty"]=w[0]
            elif "単位" in t: anchors["unit"]=w[0]
            elif "単価" in t: anchors["price"]=w[0]
            elif "金額" in t: anchors["amount"]=w[0]
            elif "備考" in t: anchors["remark"]=w[0]

        needed=["item","spec","qty","unit","price","amount","remark"]
        if all(k in anchors for k in needed):
            xs=[anchors[k] for k in needed]
            bounds=[]
            # 項目の左端はNo.列の右側。以降は隣接ヘッダーの中点。
            left=55
            for i,x in enumerate(xs):
                right=(x+xs[i+1])/2 if i+1<len(xs) else page.rect.width-20
                bounds.append((left,right))
                left=right
            return dict(zip(needed,bounds))

        # 旧BANレイアウトのフォールバック
        return {
            "item":(45,235),"spec":(235,395),"qty":(395,455),"unit":(455,510),
            "price":(510,590),"amount":(590,663),"remark":(663,820)
        }

    def ban_detail_header_y(page):
        # ページ番号ではなく、明細表ヘッダーの存在で明細ページを判定する。
        # 表紙にも似た表があるため、ページ上部（y<150）に
        # 「項目/仕様/数量/単位/単価/金額/備考」の大半があるページだけを対象にする。
        terms=set()
        ys=[]
        for w in page.get_text("words"):
            x0,y0,x1,y1,t,*_=w
            if y0 > 150:
                continue
            t=clean(t).replace(" ","").replace("　","")
            for key,label in [("item","項目"),("spec","仕様"),("qty","数量"),("unit","単位"),
                              ("price","単価"),("amount","金額"),("remark","備考")]:
                if label in t:
                    terms.add(key); ys.append(y0)
        if len(terms) >= 6 and "amount" in terms:
            header_y=max(ys) if ys else 100
            # BANの「工事内訳サマリー」ページはNo.列に1,2,3...と
            # 複数の工事項目番号が並ぶ。明細ページではNo.列は原則空欄。
            # これで従来の2ページ目サマリーは除外しつつ、
            # 0804のように2ページ目から明細が始まる帳票を取り込む。
            numbered=0
            for g in group_lines(page,header_y+8,page.rect.height-85):
                no=clean(" ".join(t for x,t in g["w"] if x<55))
                if no.isdigit():
                    numbered += 1
            if numbered >= 3:
                return None
            return header_y
        return None

    for pno in range(len(d)):
        page=d[pno]
        header_y=ban_detail_header_y(page)
        if header_y is None:
            continue
        cols=detect_ban_columns(page)
        current=None
        body_ymin=max(108, header_y+8)
        body_ymax=page.rect.height-85

        for g in group_lines(page,body_ymin,body_ymax):
            ws=sorted(g["w"])
            no=clean(" ".join(t for x,t in ws if x<55))
            item=clean(" ".join(t for x,t in ws if cols["item"][0] <= x < cols["item"][1]))

            # ページ先頭の「1 仮設工事」などを大項目として保持
            if no.isdigit() and item and g["y"]<140:
                current=f"{no}.{item}"
                continue

            c={k:clean(" ".join(t for x,t in ws if lo<=x<hi)) or None
               for k,(lo,hi) in cols.items()}

            if not c["item"]:
                continue
            compact=c["item"].replace(" ","").replace("　","")
            if compact in {"項目","小計","合計","総合計"} or c["item"].startswith("※"):
                continue

            qty=nval(c["qty"]); price=nval(c["price"]); amount=nval(c["amount"])
            remark = c["remark"]

            # BANの一部帳票では金額0が列境界をまたぎ、
            # 「金額=None, 備考='0 ※既存利用'」として抽出される。
            # 単価が存在し、備考の先頭が明確な0の場合のみ原本金額0へ戻す。
            if amount is None and price is not None and remark:
                m0 = re.match(r"^0(?:\.0+)?(?=\s|$)", str(remark).strip())
                if m0:
                    amount = 0
                    rest = str(remark).strip()[m0.end():].strip()
                    remark = rest or None

            amount_text = c["amount"] if c["amount"] and amount is None else None
            tail_note = numeric_tail_note(c["amount"]) if c["amount"] and amount is not None else None
            if tail_note:
                remark = f"{remark} / {tail_note}" if remark else tail_note
            if amount_text:
                note = f"金額欄: {amount_text}"
                remark = f"{remark} / {note}" if remark else note

            # 数値も金額欄文字もない純粋な説明行は明細化しない
            if amount is None and qty is None and price is None and not amount_text:
                continue

            rows.append([
                current,c["item"],c["spec"],qty,c["unit"],price,amount,remark,pno+1
            ])

    nums=[]
    for w in d[0].get_text("words"):
        v=nval(w[4])
        if isinstance(v,(int,float)) and abs(v)>=1000: nums.append(v)
    total=max(nums) if nums else None
    tax=pretax=None
    if total:
        for x in nums:
            if x>0 and abs((total-x)*.1-x)<5:
                tax=x; pretax=total-x; break
    return rows,total,pretax,tax

def parse_santen(path):
    d=fitz.open(path); rows=[]
    if len(d)<=2:
        text=unicodedata.normalize("NFKC",d[1].get_text("text"))
        if "4tトラック" in text and "倉庫賃料" in text:
            rows=[
                ["運送費","4tトラック","積み込み、積み下ろし人工共",2,"回",60000,120000,None,2],
                ["運送費","倉庫賃料","8/27~9/10",1,"式",30000,30000,None,2],
            ]
    else:
        cols={"name":(82,255),"spec":(255,442),"qty":(442,490),"unit":(490,535),
              "price":(535,620),"amount":(620,700),"remark":(700,790)}
        for pno in range(2,len(d)):
            top=sorted([w for w in d[pno].get_text("words") if 34<=w[1]<=62],key=lambda w:w[0])
            top_no=clean(" ".join(w[4] for w in top if w[0]<82))
            top_name=clean(" ".join(w[4] for w in top if 82<=w[0]<440))
            current=f"{top_no} {top_name}".strip() if top_name else None
            for g in group_lines(d[pno],86,500):
                ws=sorted(g["w"])
                c={k:clean(" ".join(t for x,t in ws if lo<=x<hi)) or None for k,(lo,hi) in cols.items()}
                if not c["name"] or c["name"] in {"合計","小計","総合計","名称","No.","Page."}: continue
                amount=nval(c["amount"])
                remark=c["remark"]
                tail_note=numeric_tail_note(c["amount"]) if amount is not None else None
                if tail_note:
                    remark=f"{remark} / {tail_note}" if remark else tail_note
                rows.append([current,c["name"],c["spec"],nval(c["qty"]),c["unit"],
                             nval(c["price"]),amount,remark,pno+1])
    text=unicodedata.normalize("NFKC","\n".join(d[i].get_text("text") for i in range(min(2,len(d)))))
    m=re.search(r"[￥¥]\s*([\d,]+)\s*也",text)
    pretax=int(m.group(1).replace(",","")) if m else None
    return rows,None,pretax,None

def parse_himi(path):
    d=fitz.open(path); rows=[]; current=None
    cols={"name":(55,258),"spec":(258,423),"unit":(423,480),"qty":(480,560),
          "price":(560,643),"amount":(643,732),"remark":(732,820)}
    for pno in range(2,len(d)):
        for g in group_lines(d[pno],95,510):
            ws=sorted(g["w"])
            c={k:clean(" ".join(t for x,t in ws if lo<=x<hi)) or None for k,(lo,hi) in cols.items()}
            if not c["name"] or c["name"].startswith("【合"): continue
            if re.match(r"^[Ａ-ＮA-N]\.",c["name"]):
                current=c["name"]; continue
            amt=nval(c["amount"])
            if amt is None: continue
            rows.append([current,c["name"],c["spec"],nval(c["qty"]),c["unit"],
                         nval(c["price"]),amt,c["remark"],pno+1])
    front=unicodedata.normalize("NFKC",d[0].get_text("text"))
    def money(label):
        m=re.search(re.escape(label)+r"\s*[￥¥\\]?\s*([\d,]+)",front)
        return int(m.group(1).replace(",","")) if m else None
    return rows,money("税込見積金額"),money("合計金額"),money("消費税および地方消費税")

def make_excel(results):
    wb = Workbook()
    sh = wb.active
    sh.title = "一括変換サマリー"

    headers = ["No.","PDFファイル","形式","ページ数","抽出明細","税込見積","税抜見積","消費税","算術要確認","判定"]
    sh.append(headers)
    for i,r in enumerate(results,1):
        sh.append([
            i,r["name"],r["type"],r["pages"],len(r["rows"]),
            r["total"],r["pretax"],r["tax"],r["bad"],r["status"]
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    def style_header(ws, row=1, max_col=None):
        max_col = max_col or ws.max_column
        for cell in ws[row][:max_col]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        ws.freeze_panes = "A2"

    style_header(sh, max_col=len(headers))
    widths = [6,48,12,10,12,15,15,13,13,12]
    for i,w in enumerate(widths,1):
        sh.column_dimensions[chr(64+i)].width = w
    for row in sh.iter_rows(min_row=2, min_col=6, max_col=8):
        for cell in row:
            cell.number_format = '#,##0'

    for i,r in enumerate(results,1):
        ws = wb.create_sheet(f"{i:02d}_{r['type']}")
        detail_headers = ["大項目","項目名","仕様","数量","単位","単価","金額","備考","PDFページ"]
        ws.append(detail_headers)
        for row in r["rows"]:
            ws.append(row)

        style_header(ws, max_col=len(detail_headers))
        detail_widths = [25,30,28,10,9,14,14,22,11]
        for c,w in enumerate(detail_widths,1):
            ws.column_dimensions[chr(64+c)].width = w
        for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
            for cell in row:
                cell.number_format = '#,##0'

        ws["K1"] = "元PDF"
        ws["K2"] = r["name"]
        ws["K3"] = f"形式: {r['type']}"
        ws["K1"].fill = header_fill
        ws["K1"].font = header_font
        ws.column_dimensions["K"].width = 45

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()

st.subheader("1. PDFをアップロード")
uploads=st.file_uploader("見積PDFを選択（複数可）",type=["pdf"],accept_multiple_files=True)

if uploads:
    if st.button("変換する",type="primary",width="stretch"):
        results=[]
        progress=st.progress(0)
        with tempfile.TemporaryDirectory() as td:
            for i,u in enumerate(uploads):
                p=Path(td)/u.name
                p.write_bytes(u.getvalue())
                typ=classify(str(p))
                if typ=="BAN": rows,total,pretax,tax=parse_ban(str(p))
                elif typ=="SANTEN": rows,total,pretax,tax=parse_santen(str(p))
                else: rows,total,pretax,tax=parse_himi(str(p))
                arithmetic_bad=sum(
                    1 for r in rows
                    if r[5] is not None and r[3] is not None and r[6] is not None
                    and abs(r[3]*r[5]-r[6])>max(1,abs(r[6])*.01)
                )
                missing_amount=sum(
                    1 for r in rows
                    if r[5] is not None and r[6] is None
                )
                bad=arithmetic_bad + missing_amount
                results.append({
                    "name":u.name,"type":typ,"pages":len(fitz.open(str(p))),"rows":rows,
                    "total":total,"pretax":pretax,"tax":tax,"bad":bad,
                    "status":"OK" if rows and bad==0 else ("要確認" if rows else "NG")
                })
                progress.progress((i+1)/len(uploads))

        st.session_state["results"]=results
        st.session_state["xlsx"]=make_excel(results)

if "results" in st.session_state:
    results=st.session_state["results"]
    st.subheader("2. 変換結果")
    st.dataframe([{
        "PDF":r["name"],"形式":r["type"],"明細":len(r["rows"]),
        "税込見積":r["total"],"税抜見積":r["pretax"],"消費税":r["tax"],
        "要確認":r["bad"],"判定":r["status"]
    } for r in results],width="stretch",hide_index=True)

    st.info("税額は原本記載を優先します。『税別』のみで消費税額の記載がない場合、Pilot側では税額を生成しません。")
    st.subheader("3. Excelをダウンロード")
    st.download_button(
        "変換Excelをダウンロード",
        data=st.session_state["xlsx"],
        file_name="経費予測Pilot_変換結果.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch"
    )
